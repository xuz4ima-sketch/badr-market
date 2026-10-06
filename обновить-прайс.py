# Пересобирает прайс сайта из PDF-файла «Прайс на ДД.ММ.ГГ.pdf» в корне проекта.
# Если PDF несколько, берётся тот, у которого дата «Цены указаны на …» самая свежая.
# Что создаётся:
#   site/data/prices.csv        — прайс в виде таблицы (читает сайт и этот скрипт);
#   site/data/prices.js         — копия прайса, чтобы он открывался даже при запуске index.html двойным щелчком;
#   site/data/badr-price.xlsx   — прайс в Excel для кнопки «Скачать Excel».
# Нужны пакеты: pip install pymupdf openpyxl
# Запуск: python обновить-прайс.py
# Если PDF в папке нет, прайс собирается из уже лежащего site/data/prices.csv.
import csv, io, json, pathlib, re, sys

for stream in (sys.stdout, sys.stderr):
    stream.reconfigure(encoding="utf-8")  # чтобы русский текст читался в консоли Windows

base = pathlib.Path(__file__).parent
root = base / "site"
data = root / "data"


def pdf_to_rows(path):
    """Читает PDF из 1С: строка с отступом меньше 8 — ценовая группа (издательство), с отступом 8 и больше — книга,
    дальше идут оптовая цена, ед., розничная цена, ед. Возвращает (дата, [(серия, название, опт, розница, ед.)])."""
    import pymupdf
    text = "\n".join(p.get_text() for p in pymupdf.open(path))
    lines = text.split("\n")
    m = re.search(r"Цены указаны на\s+(\d{2}\.\d{2}\.\d{4})", text)
    date = m.group(1) if m else None
    skip = {"Цена", "Ед.", "Оптовая", "Розничная"}
    is_price = lambda t: re.fullmatch(r"[\d\s ]+,\d\d\s*руб\.?", t.strip()) is not None

    def num(t):
        n = float(re.sub(r"[\s ]|руб\.?", "", t).replace(",", "."))
        return int(n) if n.is_integer() else n

    i = next((k for k, l in enumerate(lines) if l.strip() == "Розничная"), -1) + 1
    series, out = None, []
    while i < len(lines):
        line = lines[i]
        t = line.strip()
        indent = len(line) - len(line.lstrip(" "))
        if not t or t in skip:
            i += 1
        elif indent < 8:
            series = t
            i += 1
        else:
            name = t
            i += 1
            while i < len(lines) and not is_price(lines[i]):
                name += " " + lines[i].strip()
                i += 1
            if i + 3 >= len(lines):
                break
            out.append((series or "", name, num(lines[i]), num(lines[i + 2]), lines[i + 1].strip() or "шт"))
            i += 4
    return date, out


def newest_pdf():
    best = None
    for path in base.glob("Прайс*.pdf"):
        date, rows = pdf_to_rows(path)
        key = tuple(reversed(date.split("."))) if date else ()
        if rows and (best is None or key > best[0]):
            best = (key, path, date, rows)
    return best


found = newest_pdf()
if found:
    _, pdf, pdf_date, pdf_rows = found
    buf = io.StringIO(newline="")
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(["Серия", "Наименование", "Опт", "Розница", "Ед.", "Наличие", "Новинка"])
    for series, name, opt, ret, unit in pdf_rows:
        w.writerow([series, name, opt, ret, unit, "", ""])
    if pdf_date:
        w.writerow(["#обновлено", pdf_date, "", "", "", "", ""])
    (data / "prices.csv").write_text(buf.getvalue(), encoding="utf-8")
    print(f"Прайс из {pdf.name}: {len(pdf_rows)} позиций, цены на {pdf_date}")
else:
    print("PDF «Прайс*.pdf» не найден, прайс собирается из site/data/prices.csv", file=sys.stderr)

text = (data / "prices.csv").read_text(encoding="utf-8-sig")
(data / "prices.js").write_text(
    "// Создаётся автоматически из prices.csv скриптом обновить-прайс.py. Не редактируйте вручную.\n"
    "window.BADR_PRICES_CSV = " + json.dumps(text, ensure_ascii=False) + ";\n",
    encoding="utf-8")
print("Готово:", data / "prices.js")

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

NAVY, MOON = "10142B", "EFDFB0"


def norm(s):
    return s.strip().lower().replace("ё", "е")


def price(raw):
    t = re.sub(r"[\s  ]", "", raw or "")
    t = re.sub(r"(руб\.?|р\.|₽)$", "", t, flags=re.I)
    if not t:
        return None
    t = t.replace(",", "") if re.fullmatch(r"\d{1,3}(,\d{3})+(\.\d+)?", t) else t.replace(",", ".")
    try:
        n = float(t)
    except ValueError:
        return None
    return (int(n) if n.is_integer() else n) if n > 0 else None


# Same reading rules as the site (js/app.js → buildData).
rows = list(csv.reader(io.StringIO(text)))
head = [norm(h) for h in rows[0]]
col = lambda name: head.index(name) if name in head else -1
i_series = col("серия") if col("серия") >= 0 else 0
i_name, i_opt, i_ret = col("наименование"), col("опт"), col("розница")
i_unit, i_stock, i_new = col("ед."), col("наличие"), col("новинка")
groups, date = {}, None
for r in rows[1:]:
    get = lambda i: r[i].strip() if 0 <= i < len(r) else ""
    first = (r[0] if r else "").strip()
    if first.startswith("#"):
        if norm(first).startswith("#обновлено"):
            date = get(1) or None
        continue
    name = get(i_name)
    if not name:
        continue
    stock = norm(get(i_stock))
    notes = []
    if get(i_new):
        notes.append("Новинка")
    if stock.startswith("под"):
        notes.append("Под заказ")
    elif stock == "нет":
        notes.append("Нет в наличии")
    groups.setdefault(get(i_series) or "Другое", []).append(
        (re.sub(r'"([^"]*)"', r"«\1»", name), price(get(i_opt)), price(get(i_ret)), get(i_unit) or "шт", ", ".join(notes)))

config = (root / "js" / "config.js").read_text(encoding="utf-8")
phone = re.search(r'phone:\s*"([^"]*)"', config)
phone = phone.group(1).strip() if phone else ""

wb = Workbook()
ws = wb.active
ws.title = "Прайс"
wb.properties.title = "Прайс-лист Издательского Дома «BadrBook»"
wb.properties.creator = "Издательский Дом «BadrBook»"

ws.append(["Издательский Дом «BadrBook» — прайс-лист"])
ws["A1"].font = Font(bold=True, size=16, color=NAVY)
info = []
if date:
    info.append(f"Цены актуальны на {date}")
if phone:
    info.append(f"Телефон, WhatsApp и Telegram: {phone}")
ws.append(["   ·   ".join(info)])
ws["A2"].font = Font(size=11, color="595959")
ws.append([])

header = ["№", "Наименование", "Опт, ₽", "Розница, ₽", "Ед.", "Примечание"]
ws.append(header)
HEAD_ROW = ws.max_row
thin = Side(style="thin", color="DADADA")
for c in ws[HEAD_ROW]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(horizontal="center", vertical="center")

n = 0
for series, items in groups.items():
    ws.append([series])
    r = ws.max_row
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(header))
    ws.cell(r, 1).font = Font(bold=True, size=12, color=NAVY)
    ws.cell(r, 1).fill = PatternFill("solid", fgColor=MOON)
    for name, opt, ret, unit, note in items:
        n += 1
        ws.append([n, name, opt, ret, unit, note])
        r = ws.max_row
        for c in ws[r]:
            c.border = Border(bottom=thin)
        ws.cell(r, 1).alignment = Alignment(horizontal="center")
        ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="center")
        ws.cell(r, 3).font = Font(bold=True)
        for ci in (3, 4):
            v = ws.cell(r, ci).value
            ws.cell(r, ci).number_format = "#,##0" if isinstance(v, int) else "#,##0.00"
        ws.cell(r, 5).alignment = Alignment(horizontal="center")

for letter, width in zip("ABCDEF", (6, 62, 12, 12, 7, 18)):
    ws.column_dimensions[letter].width = width
ws.freeze_panes = ws.cell(HEAD_ROW + 1, 1)
ws.print_title_rows = f"{HEAD_ROW}:{HEAD_ROW}"
ws.page_setup.orientation = "portrait"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True

wb.save(data / "badr-price.xlsx")
print("Готово:", data / "badr-price.xlsx", f"({n} позиций)")
